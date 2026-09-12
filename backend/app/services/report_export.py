"""Fast Excel/PDF export helpers - stream bytes, no temp files."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def rows_to_xlsx(title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Report")[:31]
    header_font = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="EA580C")
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    if not rows:
        ws.append(["No data for selected filters"])
    for r in rows:
        ws.append([("" if v is None else str(v)) for v in r])
    for col in ws.columns:
        maxlen = 10
        letter = col[0].column_letter
        for c in col:
            maxlen = max(maxlen, min(48, len(str(c.value or ""))))
        ws.column_dimensions[letter].width = maxlen + 2
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rows_to_pdf(title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]], meta: str = "") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"<b>{title}</b>", styles["Title"]),
        Paragraph(meta or datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"), styles["Normal"]),
        Spacer(1, 12),
    ]
    data = [list(headers)] + [[("" if v is None else str(v)) for v in r] for r in rows[:2000]]
    if len(data) == 1:
        data.append(["-"] * len(headers))
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EA580C")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.Color(1, 0.97, 0.94)]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buf.getvalue()
